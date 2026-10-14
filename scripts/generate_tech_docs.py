# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x41, 0x55) # Slate 700

    # Colors
    PRIMARY = RGBColor(0x02, 0x84, 0xC7) # Sky 600
    DARK = RGBColor(0x0F, 0x17, 0x2A) # Slate 900
    MUTED = RGBColor(0x64, 0x74, 0x8B) # Slate 500

    # Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = title_p.add_run("DỰ ÁN NGHIÊN CỨU KHOA HỌC KỸ THUẬT (KHKT)\n")
    run_sub.font.size = Pt(10)
    run_sub.font.bold = True
    run_sub.font.color.rgb = PRIMARY

    run_title = title_p.add_run("BÁO CÁO CÔNG NGHỆ & KIẾN TRÚC HỆ THỐNG\nDỰ ÁN CYBERSAFE QUEST")
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = DARK

    # Subtitle info box
    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_p.paragraph_format.space_after = Pt(18)
    r_meta = info_p.add_run("Nền Tảng Giáo Dục & Rèn Luyện Phản Xạ Phòng Ngừa Lừa Đảo Trực Tuyến\nKhẩu hiệu: 1. Dừng lại  ➔  2. Kiểm tra  ➔  3. Xác minh  ➔  4. Mới hành động\nPhiên bản: 1.0.0 (Release)")
    r_meta.font.size = Pt(10)
    r_meta.font.italic = True
    r_meta.font.color.rgb = MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 1. TỔNG QUAN
    h1 = doc.add_heading(level=1)
    r = h1.add_run("1. TỔNG QUAN DỰ ÁN & MỤC TIÊU KHOA HỌC")
    r.font.color.rgb = PRIMARY
    r.font.bold = True

    p = doc.add_paragraph()
    p.add_run("CyberSafe Quest ").bold = True
    p.add_run("là một nền tảng web ứng dụng khoa học công nghệ hiện đại nhằm giải quyết vấn nạn lừa đảo trực tuyến đang bùng nổ tại Việt Nam. Nền tảng được thiết kế chuyên biệt phục vụ đề tài nghiên cứu Khoa học Kỹ thuật (KHKT), hướng đến hai nhóm đối tượng dễ bị tổn thương nhất trên không gian mạng là ")
    p.add_run("học sinh, sinh viên").bold = True
    p.add_run(" và ")
    p.add_run("người cao tuổi").bold = True
    p.add_run(".\n\nNền tảng kết hợp phương pháp giáo dục hiện đại:")
    
    b1 = doc.add_paragraph(style='List Bullet')
    b1.add_run("Triết lý giáo dục: ").bold = True
    b1.add_run("Đồng cảm & Không đổ lỗi nạn nhân, giúp người học nhận biết các kỹ thuật thao túng tâm lý (Social Engineering).")

    b2 = doc.add_paragraph(style='List Bullet')
    b2.add_run("Quy trình 4 bước chuẩn quốc gia: ").bold = True
    b2.add_run("Dừng lại ➔ Kiểm tra ➔ Xác minh ➔ Mới hành động.")

    b3 = doc.add_paragraph(style='List Bullet')
    b3.add_run("Đánh giá 4 trụ cột năng lực số: ").bold = True
    b3.add_run("Nhận diện (Detection), Kiểm chứng (Verification), Bảo vệ tài khoản (Account Protection), Ứng phó sự cố (Incident Response).")

    # 2. BẢNG TỔNG HỢP CÔNG NGHỆ
    h2 = doc.add_heading(level=1)
    r = h2.add_run("2. BẢNG TỔNG HỢP DANH MỤC CÔNG NGHỆ SỬ DỤNG")
    r.font.color.rgb = PRIMARY
    r.font.bold = True

    table_data = [
        ["STT", "Công nghệ / Thư viện", "Phiên bản", "Phân loại", "Vai trò & Chức năng trong hệ thống"],
        ["1", "Next.js", "^15.1.7", "Fullstack Framework", "Framework chính, App Router, SSR, Server/Client Components, định tuyến động"],
        ["2", "React", "^19.0.0", "Core UI Library", "Xây dựng giao diện tương tác, quản lý State và Lifecycle của trò chơi giải đố"],
        ["3", "TypeScript", "^5.7.3", "Programming Language", "Ngôn ngữ phát triển chính, định kiểu tĩnh nghiêm ngặt (Strict Types) cho 39 ải"],
        ["4", "Tailwind CSS", "^4.0.0", "CSS Framework", "Hệ thống Styling utility-first, giao diện sáng (Light Mode), tối ưu responsive"],
        ["5", "Supabase", "^2.48.1", "Backend as a Service", "Cơ sở dữ liệu PostgreSQL Cloud, quản lý kịch bản, xác thực Admin, bảo mật RLS"],
        ["6", "Lucide React", "^0.475.0", "Iconography", "Bộ icon vector hiện đại, trực quan hóa 11 nhóm chiêu thức và cảnh báo rủi ro"],
        ["7", "Framer Motion", "^12.4.7", "Animation Library", "Xử lý hiệu ứng chuyển động mượt mà cho câu hỏi, modal và thông báo kết quả"],
        ["8", "Canvas Confetti", "^1.9.4", "Visual FX", "Hiệu ứng pháo giấy chúc mừng khi hoàn thành xuất sắc ải hoặc đạt 3 sao"],
        ["9", "clsx & tailwind-merge", "^2.1.1 / ^3.0.1", "CSS Utilities", "Gộp và chuẩn hóa className điều kiện linh hoạt không bị xung đột class"],
        ["10", "SVG Math Vector", "Native SVG", "Vector Graphics", "Thuật toán đường cong Bezier vẽ lộ trình hành trình tự co giãn theo chiều cao"]
    ]

    table = doc.add_table(rows=len(table_data), cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(0.5), Inches(1.5), Inches(0.9), Inches(1.4), Inches(2.2)]

    for row_idx, row in enumerate(table.rows):
        is_header = (row_idx == 0)
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths[col_idx]
            cell.text = table_data[row_idx][col_idx]
            set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.runs[0]
            
            if is_header:
                set_cell_background(cell, "0284C7") # Sky 600
                run.font.bold = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                run.font.size = Pt(9)
                if col_idx in [0, 2]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if row_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC") # Slate 50
                else:
                    set_cell_background(cell, "FFFFFF")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 3. CHI TIẾT CÔNG NGHỆ CỐT LÕI
    h3 = doc.add_heading(level=1)
    r = h3.add_run("3. PHÂN TÍCH CHI TIẾT CÔNG NGHỆ CỐT LÕI")
    r.font.color.rgb = PRIMARY
    r.font.bold = True

    # 3.1 Next.js 15
    doc.add_heading("3.1. Next.js 15 (App Router) & React 19", level=2)
    p = doc.add_paragraph()
    p.add_run("Next.js 15 mang lại khả năng kết hợp tối ưu giữa Server-Side Rendering (SSR) và Client Components:\n")
    p.add_run("• Server Components: ").bold = True
    p.add_run("Được sử dụng cho trang chủ và các trang thông tin tĩnh giúp nạp trang tức thì, thân thiện với SEO và tối ưu hóa tốc độ ban đầu (FCP).\n")
    p.add_run("• Dynamic Routing (/play/[chapterId]): ").bold = True
    p.add_run("Định tuyến linh hoạt cho 11 đề tài chuyên sâu, mỗi đề tài có URL độc lập, cho phép người học lưu và chia sẻ tiến độ học tập cụ thể.\n")
    p.add_run("• React 19 State Management: ").bold = True
    p.add_run("Sử dụng custom hook (usePlayerJourney) để theo dõi trạng thái mở khóa ải, điểm số, chuỗi ngày rèn luyện và tiến trình tuần tự.")

    # 3.2 TypeScript
    doc.add_heading("3.2. TypeScript 5.7 (Kiểm Soát Kiểu Dữ Liệu Chặt Chẽ)", level=2)
    p = doc.add_paragraph()
    p.add_run("Toàn bộ mã nguồn được bao bọc bởi hệ thống kiểu dữ liệu TypeScript nghiêm ngặt:\n")
    p.add_run("• Định nghĩa Scenario: ").bold = True
    p.add_run("Quy định chuẩn cấu trúc câu hỏi, loại thiết bị mô phỏng (call, sms, social, url), danh sách lựa chọn (choices) và giải thích giáo dục.\n")
    p.add_run("• Ngăn ngừa lỗi Runtime: ").bold = True
    p.add_run("Đảm bảo các thao tác giải đố, tính điểm sao, tính % hoàn thành luôn hoạt động chính xác tuyệt đối không bị lỗi crash.")

    # 3.3 Tailwind CSS v4 & Thiết kế Thân Thiện
    doc.add_heading("3.3. Tailwind CSS v4 & Trải Nghiệm Giao Diện Người Dùng (UI/UX)", level=2)
    p = doc.add_paragraph()
    p.add_run("Khác với các giao diện tối màu phức tạp, CyberSafe Quest áp dụng phong cách thiết kế chuyên biệt:\n")
    p.add_run("• Thiết kế Thân thiện Đa Thế Hệ (Inclusive Light Design): ").bold = True
    p.add_run("Nền trắng sáng, độ tương phản chữ cao (Slate 900 trên nền trắng/Slate 50), phông chữ to rõ ràng, hỗ trợ tối đa cho người cao tuổi và các em học sinh.\n")
    p.add_run("• Hiển thị Trọn vẹn Tranh Minh Họa (Full Uncropped Art): ").bold = True
    p.add_run("Toàn bộ 13 bức tranh vẽ tay độc quyền được hiển thị theo tỷ lệ gốc (object-contain), không bị cắt xén, tách biệt hoàn toàn huy hiệu ra khỏi tranh để người dùng quan sát rõ từng chi tiết cảnh báo.\n")
    p.add_run("• Responsive 100%: ").bold = True
    p.add_run("Tương thích hoàn hảo từ màn hình điện thoại di động (Mobile-first) đến máy tính bảng và màn hình máy tính bàn.")

    # 3.4 Thuật toán SVG Bezier
    doc.add_heading("3.4. Thuật Toán Đồ Họa Vector Toán Học (Bản Đồ Hành Trình SVG)", level=2)
    p = doc.add_paragraph()
    p.add_run("Bản đồ hành trình rèn luyện được vẽ bằng mã toán học SVG thời gian thực:\n")
    p.add_run("• Đường cong Bezier bậc ba: ").bold = True
    p.add_run("Sử dụng lệnh M (Move) và C (Cubic Bezier Curve) nối giữa các tọa độ trạm ải (25% chiều ngang cho ải bên trái và 75% cho ải bên phải).\n")
    p.add_run("• Hệ quy chiếu co giãn linh hoạt: ").bold = True
    p.add_run("Sử dụng viewBox='0 0 1000 ${height}' kết hợp preserveAspectRatio='none', giúp vệt đường nhựa và vạch kẻ tim đường màu trắng luôn kết nối chính xác vào tâm của từng trạm ải dù danh sách có 3 hay 30 câu hỏi.")

    # 3.5 Supabase & Hybrid Storage
    doc.add_heading("3.5. Cơ Sở Dữ Liệu Supabase & Mô Hình Lưu Trữ Kép (Hybrid Offline-Online)", level=2)
    p = doc.add_paragraph()
    p.add_run("Hệ thống giải quyết bài toán môi trường mạng không ổn định tại trường học hoặc vùng sâu vùng xa:\n")
    p.add_run("• Cloud Database: ").bold = True
    p.add_run("Supabase PostgreSQL lưu trữ 39 kịch bản kiểm chuẩn, bảng ghi lịch sử làm bài (user_progress), hồ sơ học viên (user_profiles), và lịch sử tra cứu tin nhắn (chatbot_consultations).\n")
    p.add_run("• Bảo mật Row Level Security (RLS): ").bold = True
    p.add_run("Học viên chỉ có quyền đọc kịch bản công khai; chỉ tài khoản có cờ is_admin = true mới có quyền thêm/chỉnh sửa nội dung câu hỏi trong database.\n")
    p.add_run("• Offline Fallback (LocalStorage): ").bold = True
    p.add_run("Khi không có kết nối Internet, hệ thống tự động nạp bộ 39 kịch bản tĩnh có sẵn và lưu kết quả vượt ải vào LocalStorage của trình duyệt. Khi có mạng trở lại, dữ liệu được tự động đồng bộ lên đám mây.")

    # 4. Ý NGHĨA KHOA HỌC KỸ THUẬT
    h4 = doc.add_heading(level=1)
    r = h4.add_run("4. Ý NGHĨA ĐÓNG GÓP CHO ĐỀ TÀI NGHIÊN CỨU KHOA HỌC (KHKT)")
    r.font.color.rgb = PRIMARY
    r.font.bold = True

    p = doc.add_paragraph()
    p.add_run("Hệ thống công nghệ được phát triển không chỉ là một ứng dụng web thông thường mà đóng vai trò là một ")
    p.add_run("công cụ thực nghiệm đo lường khoa học:").bold = True

    b1 = doc.add_paragraph(style='List Bullet')
    b1.add_run("Thu thập dữ liệu thực nghiệm định lượng: ").bold = True
    b1.add_run("Ghi nhận số lần thử (attempts), điểm số (score), thời gian xử lý và tỷ lệ an toàn của từng đối tượng học viên theo thời gian thực.")

    b2 = doc.add_paragraph(style='List Bullet')
    b2.add_run("Trung tâm Quản trị & Xuất Báo Cáo CSV (/admin): ").bold = True
    b2.add_run("Cho phép Ban giám khảo và nhóm nghiên cứu trích xuất toàn bộ dữ liệu khảo sát ra định dạng Excel (.csv) để vẽ biểu đồ và phân tích thống kê phương sai (ANOVA, T-Test).")

    b3 = doc.add_paragraph(style='List Bullet')
    b3.add_run("Tính tuần tự bắt buộc (Gamification chuẩn hóa): ").bold = True
    b3.add_run("Người học không thể gian lận nhảy cóc (skip), bắt buộc phải vận dụng kiến thức nhận diện đúng ải hiện tại mới được mở khóa ải tiếp theo.")

    # 5. KẾT LUẬN
    h5 = doc.add_heading(level=1)
    r = h5.add_run("5. KẾT LUẬN")
    r.font.color.rgb = PRIMARY
    r.font.bold = True

    p = doc.add_paragraph()
    p.add_run("Sự kết hợp đồng bộ giữa ")
    p.add_run("Next.js 15, React 19, TypeScript, Tailwind CSS, Supabase và kiến trúc Hybrid Offline-Online").bold = True
    p.add_run(" đã tạo nên một giải pháp công nghệ toàn diện, thẩm mỹ cao, vận hành bền bỉ và đáp ứng hoàn hảo các tiêu chí khắt khe của một công trình nghiên cứu Khoa học Kỹ thuật dành cho cộng đồng.")

    # Save
    doc.save("Bao cao cong nghe su dung - CyberSafe Quest.docx")
    print("SUCCESS: Document created successfully!")

if __name__ == "__main__":
    create_document()
